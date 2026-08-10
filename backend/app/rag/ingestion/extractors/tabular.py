"""CSV and XLSX extraction.

A spreadsheet in the *knowledge base* is treated as a document to read, not as a
dataset to query — structured analysis over tabular data is the `datasets` table
and the typed analysis tools in Phase 6. The distinction matters: retrieval over a
100k-row CSV is the wrong tool for "what was Q2 revenue", and this path does not
pretend otherwise. It exists so a small reference table shipped alongside a report
is searchable.

Rows are emitted in groups with the header repeated on each block, because a bare
row of numbers is meaningless once separated from its column names — and after
chunking, separated is exactly what it would be.
"""

from __future__ import annotations

import csv
import io
from typing import Any

import openpyxl

from app.core.exceptions import DocumentProcessingError
from app.rag.ingestion.blocks import Block, BlockKind, ExtractedDocument
from app.rag.ingestion.extractors.text import decode

#: Data rows per block. Small enough that a block stays well inside the chunk
#: budget once the header is repeated on it.
ROWS_PER_BLOCK = 20
#: Hard ceiling on rows read from one sheet. A spreadsheet larger than this is a
#: dataset, and Phase 6 is where it belongs.
MAX_ROWS = 5_000
MAX_CELL_CHARS = 500


def _clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().replace("\n", " ")
    return text[:MAX_CELL_CHARS]


def _rows_to_blocks(header: list[str], rows: list[list[str]], *, sheet: str | None) -> list[Block]:
    if not rows:
        return []

    header_line = " | ".join(header) if any(header) else ""
    blocks: list[Block] = []

    for start in range(0, len(rows), ROWS_PER_BLOCK):
        window = rows[start : start + ROWS_PER_BLOCK]
        lines = [" | ".join(row) for row in window if any(cell for cell in row)]
        if not lines:
            continue
        body = "\n".join(lines)
        text = f"{header_line}\n{body}" if header_line else body
        if sheet:
            text = f"Sheet: {sheet}\n{text}"
        blocks.append(Block(text=text, kind=BlockKind.TABLE))

    return blocks


def extract_csv(payload: bytes) -> ExtractedDocument:
    text = decode(payload).replace("\r\n", "\n").replace("\r", "\n")

    try:
        # Sniffing the dialect handles semicolon-separated exports from
        # non-English Excel locales, which are common and otherwise parse as one
        # giant column.
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:8192])
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(text), dialect)
    try:
        # strict=False: the point of the zip is to stop early at MAX_ROWS, so the
        # two iterables are expected to differ in length.
        parsed = [
            [_clean(cell) for cell in row]
            for _, row in zip(range(MAX_ROWS + 1), reader, strict=False)
        ]
    except csv.Error as exc:
        raise DocumentProcessingError("The CSV file could not be parsed.") from exc

    if not parsed:
        return ExtractedDocument(blocks=[], page_count=None, title=None)

    header, rows = parsed[0], parsed[1:]
    return ExtractedDocument(
        blocks=_rows_to_blocks(header, rows, sheet=None), page_count=None, title=None
    )


def extract_xlsx(payload: bytes) -> ExtractedDocument:
    try:
        # read_only streams rows instead of building the whole object graph;
        # data_only takes cached formula results rather than "=SUM(B2:B9)".
        workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentProcessingError("The spreadsheet could not be read.") from exc

    blocks: list[Block] = []
    try:
        for sheet in workbook.worksheets:
            rows_iter = sheet.iter_rows(values_only=True, max_row=MAX_ROWS + 1)
            parsed = [[_clean(cell) for cell in row] for row in rows_iter]
            if not parsed:
                continue
            header, rows = parsed[0], parsed[1:]
            sheet_blocks = _rows_to_blocks(header, rows, sheet=str(sheet.title))
            if sheet_blocks:
                blocks.append(Block(text=str(sheet.title), kind=BlockKind.HEADING, level=1))
                blocks.extend(sheet_blocks)
    finally:
        workbook.close()

    return ExtractedDocument(blocks=blocks, page_count=None, title=None)
